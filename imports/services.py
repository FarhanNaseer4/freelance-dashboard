import csv
import io

from openpyxl import load_workbook

from clients.models import Client
from projects.models import Project


CLIENT_HEADERS = {"name", "platform", "email", "company", "notes", "first_contact_date", "last_contact_date"}
PROJECT_HEADERS = {
    "client",
    "project",
    "platform",
    "contract_type",
    "project_type",
    "start_date",
    "delivery_date",
    "budget",
    "status",
    "priority",
    "next_action",
    "follow_up_date",
}


def parse_upload(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".csv"):
        text = uploaded_file.read().decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if name.endswith(".xlsx"):
        workbook = load_workbook(uploaded_file, read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() for value in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]
    raise ValueError("Upload a CSV or XLSX file.")


def validate_and_import(rows, commit=False):
    results = []
    for index, raw in enumerate(rows, start=2):
        row = {str(key).strip().lower(): (value if value is not None else "") for key, value in raw.items()}
        errors = []
        kind = "client" if "name" in row and "project" not in row else "project"
        if kind == "client":
            if not row.get("name"):
                errors.append("Missing client name.")
            if row.get("platform") not in dict(Client.Platform.choices):
                errors.append("Invalid platform.")
            if commit and not errors:
                Client.objects.update_or_create(
                    name=row["name"],
                    defaults={
                        "platform": row.get("platform") or Client.Platform.DIRECT,
                        "email": row.get("email", ""),
                        "company": row.get("company", ""),
                        "notes": row.get("notes", ""),
                    },
                )
        else:
            if not row.get("client"):
                errors.append("Missing client.")
            if not row.get("project"):
                errors.append("Missing project.")
            if row.get("platform") not in dict(Project.Platform.choices):
                errors.append("Invalid platform.")
            if row.get("contract_type") not in dict(Project.ContractType.choices):
                errors.append("Invalid contract type.")
            if commit and not errors:
                client, _ = Client.objects.get_or_create(name=row["client"], defaults={"platform": row.get("platform") or Client.Platform.DIRECT})
                Project.objects.update_or_create(
                    client=client,
                    name=row["project"],
                    defaults={
                        "platform": row.get("platform") or Project.Platform.DIRECT,
                        "contract_type": row.get("contract_type") or Project.ContractType.FIXED,
                        "project_type": row.get("project_type") or Project.ProjectType.OTHER,
                        "budget": row.get("budget") or 0,
                        "status": row.get("status") or Project.Status.RUNNING,
                        "priority": row.get("priority") or Project.Priority.MEDIUM,
                        "next_action": row.get("next_action", ""),
                    },
                )
        results.append({"row": index, "kind": kind, "valid": not errors, "errors": errors, "data": row})
    return results
